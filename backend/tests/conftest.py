from __future__ import annotations

import io
import os
import uuid
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from gawe.api.app import create_app
from gawe.config import DatabaseConfig, load_config
from gawe.db import repo

# One schema per test run. The prefix is required for any destructive cleanup.
TEST_SCHEMA = f"{repo.TEST_SCHEMA_PREFIX}{uuid.uuid4().hex[:12]}"

CONFIG_TEMPLATE = f"""
[server]
host = "127.0.0.1"
port = 8000

[storage]
data_dir = "data"
profile_dir = "profile"
logs_dir = "logs"

[database]
host = "127.0.0.1"
port = 3306
user = "root"
name = "{TEST_SCHEMA}"

[llm]
active_provider = "provider_a"
default_timeout_s = 90
max_retries = 1
repair_json_once = true

[llm.providers.provider_a]
type = "openai_compatible"
base_url = "https://provider-a.example/v1"
api_key_ref = "gawe/llm/provider_a"
capabilities = ["chat"]

[llm.roles]
extract = "model-x"
classify = "model-x"
write = "model-y"
review = "model-z"

[privacy]
cloud_consent_required = true
log_document_content = false
log_model_io = false

[retention]
job_cache_days = 90
log_days = 30
diagnostics_days = 7

[frontend]
dist_dir = "frontend/dist"
"""

SETUP_CONFIG = CONFIG_TEMPLATE.replace(
    'active_provider = "provider_a"',
    'active_provider = "fake"',
).replace(
    '[llm.providers.provider_a]\ntype = "openai_compatible"\n'
    'base_url = "https://provider-a.example/v1"\n'
    'api_key_ref = "gawe/llm/provider_a"\n'
    'capabilities = ["chat"]',
    '[llm.providers.fake]\ntype = "fake"',
)

FIXTURE_CV = (
    Path(__file__).resolve().parents[2]
    / "docs"
    / "phase-00"
    / "fixtures"
    / "master_cv"
    / "candidate_001.md"
)


def database_config() -> DatabaseConfig:
    """Test database settings come from the environment, never from a file."""
    return DatabaseConfig(
        host=os.environ.get("GAWE_TEST_DB_HOST", "127.0.0.1"),
        port=int(os.environ.get("GAWE_TEST_DB_PORT", "3306")),
        user=os.environ.get("GAWE_TEST_DB_USER", "root"),
        name=TEST_SCHEMA,
    )


@pytest.fixture(scope="session")
def test_schema() -> str:
    settings = database_config()
    repo.drop_schema(settings)
    repo.create_schema(settings)
    connection = repo.connect(settings)
    try:
        repo.migrate(connection)
    finally:
        connection.close()
    yield TEST_SCHEMA
    repo.drop_schema(settings)


@pytest.fixture
def _clean_database(test_schema: str) -> None:
    """Start a database test from a known-empty schema.

    Not autouse: ``connection`` and ``setup_client`` request it, and modules
    that reach the database through ``config``/``repo_root`` opt in with
    ``pytestmark = pytest.mark.usefixtures("_clean_database")``. Tests that
    never open the database must not pay for a reset.
    """
    db = repo.connect(database_config())
    try:
        repo.reset(db)
    finally:
        db.close()


@pytest.fixture
def connection(test_schema: str, _clean_database):
    db = repo.connect(database_config())
    yield db
    db.close()


@pytest.fixture
def repo_root(tmp_path: Path) -> Path:
    (tmp_path / "config.toml").write_text(CONFIG_TEMPLATE, encoding="utf-8")
    return tmp_path


@pytest.fixture
def config(repo_root: Path):
    return load_config(repo_root / "config.toml")


# -- Phase 2A setup fixtures ------------------------------------------------
@pytest.fixture
def setup_root(tmp_path: Path) -> Path:
    (tmp_path / "config.toml").write_text(SETUP_CONFIG, encoding="utf-8")
    return tmp_path


@pytest.fixture
def setup_config(setup_root: Path):
    return load_config(setup_root / "config.toml")


@pytest.fixture
def setup_client(setup_config, _clean_database):
    with TestClient(create_app(setup_config)) as test_client:
        yield test_client


@pytest.fixture
def csrf(setup_client) -> dict[str, str]:
    session = setup_client.get("/api/session").json()
    return {"X-CSRF-Token": session["csrf_token"]}


def upload_cv(client: TestClient, headers: dict, data: bytes | None = None,
              name: str = "candidate_001.md"):
    payload = FIXTURE_CV.read_bytes() if data is None else data
    return client.post(
        "/api/setup/master-cv",
        headers=headers,
        files={"file": (name, payload, "application/octet-stream")},
    )


def extract_profile(client: TestClient, headers: dict):
    return client.post("/api/setup/master-cv/extractions", headers=headers)


def latest_profile(client: TestClient) -> dict:
    items = client.get("/api/setup/profiles").json()["items"]
    assert items, "no profile versions"
    return client.get(f"/api/setup/profiles/{items[0]['profile_version_id']}").json()


def make_pdf(lines: list[str]) -> bytes:
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "\n".join(lines), fontsize=9)
    data = document.tobytes()
    document.close()
    return data


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        added = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, value in enumerate(row):
                added.cell(row_index, column_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
