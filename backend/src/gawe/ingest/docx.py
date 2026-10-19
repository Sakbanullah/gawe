"""DOCX text extraction with paragraph/table locators (python-docx)."""

from __future__ import annotations

from .types import Block, DocumentText


def parse_docx(path: str) -> DocumentText:
    import docx

    document = docx.Document(path)
    blocks: list[Block] = []
    section: str | None = None

    for index, paragraph in enumerate(document.paragraphs):
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        if style.startswith("heading"):
            section = text.lower()
        blocks.append(
            Block(
                text=text,
                locator={
                    "page": None,
                    "section": section,
                    "block_id": f"para:{index}",
                    "char_start": None,
                    "char_end": None,
                },
            )
        )

    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                text = (cell.text or "").strip()
                if not text:
                    continue
                blocks.append(
                    Block(
                        text=text,
                        locator={
                            "page": None,
                            "section": section,
                            "block_id": f"table:{table_index}:row:{row_index}:cell:{cell_index}",
                            "char_start": None,
                            "char_end": None,
                        },
                    )
                )

    return DocumentText(blocks=blocks, page_count=None)
